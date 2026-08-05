# Format export/test-data.json into a reviewable Excel workbook.
#
# Deliberately generic: all domain knowledge lives in scripts/export-data.mjs,
# which emits sheets as {name, note, columns, rows}. This only knows how to make
# that look like something a treasury reviewer can open, filter and pivot.
#
# Run via `npm run export-data` (which runs the Node exporter first).
import json
import re
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "export" / "test-data.json"
OUT = ROOT / "export" / "Cash Forecast Grid - test data.xlsx"

# ---- palette ---------------------------------------------------------------
INK = "1F2933"
MUTED = "6B7785"
RULE = "D8DEE6"
HEAD_BG = "1F2933"
HEAD_FG = "FFFFFF"
BAND = "F5F7FA"
GOOD_BG, GOOD_FG = "E3F5EC", "0B6B45"
BAD_BG, BAD_FG = "FCE8EC", "A31432"

# Sheet order, family and tab colour all arrive on the descriptors from
# scripts/column-docs.js via the exporter, so there is nothing to keep in sync here.
SETUP_COLOUR = "1F2933"

FMT = {
    "money": "#,##0;[Red]-#,##0",
    "int": "#,##0",
    "num": "0.####",
    "pct": "0.0%",
    "date": "ddd dd mmm yyyy",
}
RIGHT = {"money", "int", "num", "pct"}

ISO = re.compile(r"^\d{4}-\d{2}-\d{2}$")
thin = Side(style="thin", color=RULE)


def as_date(v):
    if isinstance(v, str) and ISO.match(v):
        y, m, d = v.split("-")
        return date(int(y), int(m), int(d))
    return v


def note_height(text, total_width):
    """Row height for a merged, wrapped note spanning `total_width` characters."""
    chars = max(40, total_width)
    lines = max(1, -(-len(text) // chars))
    return 13 * lines + 6


def write_sheet(ws, spec, family_colour):
    cols = spec["columns"]
    rows = spec["rows"]
    n = len(cols)
    last_col = get_column_letter(n)
    total_width = sum(c.get("width") or 14 for c in cols)

    ws.sheet_properties.tabColor = family_colour

    # Title
    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]
    t.value = spec["name"]
    t.font = Font(bold=True, size=14, color=INK)
    t.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24

    # Note
    ws.merge_cells(f"A2:{last_col}2")
    nt = ws["A2"]
    nt.value = spec["note"]
    nt.font = Font(size=9.5, italic=True, color=MUTED)
    nt.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = note_height(spec["note"], total_width)

    # Header
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=3, column=i, value=c["label"])
        cell.font = Font(bold=True, size=10, color=HEAD_FG)
        cell.fill = PatternFill("solid", fgColor=HEAD_BG)
        cell.alignment = Alignment(
            horizontal="right" if c["type"] in RIGHT else "left",
            vertical="bottom",
            wrap_text=True,
        )
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(i)].width = c.get("width") or 14
    ws.row_dimensions[3].height = 28

    # Body
    band = PatternFill("solid", fgColor=BAND)
    for r, row in enumerate(rows, start=4):
        shaded = (r % 2) == 1
        for i, c in enumerate(cols, start=1):
            v = row[i - 1] if i - 1 < len(row) else None
            kind = c["type"]
            if kind == "date":
                v = as_date(v)
            elif kind == "bool":
                v = "Yes" if v else "No"
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = Font(size=10, color=INK)
            if kind in FMT:
                cell.number_format = FMT[kind]
            if kind == "bool":
                cell.alignment = Alignment(horizontal="center")
                cell.fill = PatternFill("solid", fgColor=GOOD_BG if v == "Yes" else BAD_BG)
                cell.font = Font(size=10, bold=True, color=GOOD_FG if v == "Yes" else BAD_FG)
                continue
            if kind == "text" and (c.get("width") or 14) >= 40:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if shaded:
                cell.fill = band

    if rows:
        ws.auto_filter.ref = f"A3:{last_col}{3 + len(rows)}"
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def write_contents(ws, meta, specs):
    ws.sheet_properties.tabColor = SETUP_COLOUR
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 11
    ws.column_dimensions["D"].width = 110
    ws.sheet_view.showGridLines = False

    ws["B1"] = meta["title"]
    ws["B1"].font = Font(bold=True, size=20, color=INK)
    ws.row_dimensions[1].height = 30

    ws["B2"] = (
        f'Generated from {meta["generatedFrom"]}  ·  horizon {meta["horizon"]}  ·  '
        f'{meta["entities"]} entities, {meta["grids"]} grids  ·  '
        f'{meta["checksPassed"]}/{meta["checksTotal"]} reconciliation checks pass'
    )
    ws["B2"].font = Font(size=10, color=MUTED)
    ws.row_dimensions[2].height = 16

    r = 4
    ws[f"B{r}"] = "Read this first"
    ws[f"B{r}"].font = Font(bold=True, size=12, color=INK)
    r += 1
    for note in meta["notes"]:
        ws[f"B{r}"] = "•"
        ws[f"B{r}"].font = Font(size=10, color=MUTED)
        ws[f"B{r}"].alignment = Alignment(horizontal="center", vertical="top")
        ws.merge_cells(f"C{r}:D{r}")
        ws[f"C{r}"] = note
        ws[f"C{r}"].font = Font(size=10, color=INK)
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = note_height(note, 118)
        r += 1

    r += 1
    # Specs arrive already in workbook order, so grouping is just a run-length pass
    # over the family each one declares.
    groups = []
    for spec in specs:
        fam = spec.get("family") or "Other"
        if not groups or groups[-1][0] != fam:
            groups.append((fam, spec.get("familyColour") or MUTED, []))
        groups[-1][2].append(spec)

    for family, colour, members in groups:
        ws[f"B{r}"] = family
        ws[f"B{r}"].font = Font(bold=True, size=12, color=colour)
        ws[f"B{r}"].border = Border(bottom=Side(style="medium", color=colour))
        for col in ("C", "D"):
            ws[f"{col}{r}"].border = Border(bottom=Side(style="medium", color=colour))
        ws.row_dimensions[r].height = 20
        r += 1
        for spec in members:
            name = spec["name"]
            link = ws[f"B{r}"]
            link.value = name
            link.hyperlink = f"#'{name}'!A1"
            link.font = Font(size=10, color="0078FF", underline="single")
            cnt = ws[f"C{r}"]
            cnt.value = len(spec["rows"])
            cnt.number_format = "#,##0"
            cnt.font = Font(size=10, color=MUTED)
            cnt.alignment = Alignment(horizontal="right")
            d = ws[f"D{r}"]
            d.value = spec["note"]
            d.font = Font(size=9.5, color=MUTED)
            d.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = note_height(spec["note"], 108)
            r += 1
        r += 1

    ws["C4"] = "Rows"
    ws["C4"].font = Font(bold=True, size=9, color=MUTED)
    ws["C4"].alignment = Alignment(horizontal="right")


def main():
    payload = json.loads(SRC.read_text(encoding="utf-8"))
    specs = payload["sheets"]  # already in workbook order

    wb = Workbook()
    write_contents(wb.active, payload["meta"], specs)
    wb.active.title = "Contents"

    for spec in specs:
        write_sheet(wb.create_sheet(spec["name"]), spec, spec.get("familyColour") or MUTED)

    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    size = OUT.stat().st_size / 1024 / 1024
    print(f"{len(specs) + 1} sheets -> {OUT} ({size:.1f} MB)")


if __name__ == "__main__":
    main()
